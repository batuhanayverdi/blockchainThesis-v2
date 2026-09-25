"""
Build a Hive review workbook: Model 3 vs Model 1 vs Model 4, claim by claim
===========================================================================
Unlike the YouTube benchmark reviewer, the Hive run has NO gold labels, so there
is no TP/FP/FN scoring here. This workbook is for INSPECTION and for coding a
validation sample: it lays the pipeline out so you can read each post, see what
each model decided, and why.

  Sheet "by_post"          one row per post: the Model 4 cascade label and which
                           model produced it (label_source), shown next to Model
                           3's 6-class label (+ ordinal) and Model 1's title
                           label, plus engagement, claim counts, how many Model 3
                           claims were the misinformation trigger (false-side AND
                           asserted), a flag for where the two models disagree,
                           and blank human_label / human_notes columns for coding.
  Sheet "by_claim_model3"  one row per Model 3 body claim: claim, query, 6-class
                           verdict, stance, justification, search results, and a
                           claim_role that shows whether it triggered the label
                           (trigger / quoted_false / true_side / unverified).
  Sheet "by_claim_model1"  one row per Model 1 title claim: same columns for the
                           title-as-claim model that the cascade falls back to.

Coloring (by_post, on the Model 4 binary): misinfo red, ok green, abstain grey.
Coloring (claim sheets, on claim_role): trigger red, quoted_false yellow,
true_side green, unverified grey. Sort or filter on any of these.

Reads (all local, produced by the Hive run):
  hive_outputs/model3/stage4_factcheck.json
  hive_outputs/model1/stage4_factcheck.json
  hive_outputs/model4/posts.csv
  hive_sample_master.parquet            (optional, only for engagement columns)

Run:  python build_hive_review_excel.py
Deps: pip install pandas openpyxl pyarrow
"""

import os
import json
import pandas as pd
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter
try:
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
except Exception:                                    # very old / new openpyxl
    import re as _re
    ILLEGAL_CHARACTERS_RE = _re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")


# ============================ CONFIG ========================================
M3_JSON = os.path.join(r"C:\Users\batuh\PycharmProjects\Thesis\fact_checked_hive_pipeline\4.3 model full transkript (6 verdict)", "hive_outputs", "model3", "stage4_factcheck.json")
M1_JSON = os.path.join(r"C:\Users\batuh\PycharmProjects\Thesis\fact_checked_hive_pipeline\4.1-model (just title)", "hive_outputs", "model1", "stage4_factcheck.json")
M4_CSV = os.path.join("hive_outputs", "model4", "posts.csv")
PARQUET = r"C:\Users\batuh\PycharmProjects\Thesis\fact_checked_hive_pipeline\hive sample preparation (politic-non-politic)\hive_sample_master.parquet"         # optional; for engagement columns
OUTPUT_XLSX = "hive_review.xlsx"

BODY_CHARS = 2000               # snippet length (Excel cells cap near 32k chars)
RESULTS_CHARS = 8000

FALSE_SIDE = {"pants_on_fire", "false", "mostly_false"}
TRUE_SIDE = {"mostly_true", "true"}

# engagement columns to pull from the parquet if present
ENGAGEMENT_COLS = ["net_votes", "children", "payout_combined"]

BINARY_ORDER = {"misinfo": 0, "ok": 1, "abstain": 2}

WRAP_COLS = {"title", "body_snippet", "results", "justification", "claim",
             "query", "model3_verdict_counts", "human_notes", "model1_query"}
WIDTHS = {"title": 40, "body_snippet": 70, "results": 70, "justification": 50,
          "claim": 50, "query": 38, "model1_query": 38, "channel": 22,
          "author": 20, "date": 20, "category": 18, "sample_stratum": 14,
          "model3_verdict_counts": 24, "verdict": 13, "stance": 20,
          "model4_label": 16, "model4_binary": 13, "label_source": 24,
          "model3_label": 15, "model3_ordinal": 13, "model1_label": 16,
          "model1_verdict": 14, "model1_stance": 20, "claim_role": 14,
          "human_label": 16, "human_notes": 40, "models_disagree": 15,
          "n_model3_claims": 14, "asserted_falseside_n": 18, "is_political": 12,
          "net_votes": 11, "children": 10, "payout_combined": 15, "doc_id": 30}
POST_FILLS = {"misinfo": "FFC7CE", "ok": "C6EFCE", "abstain": "D9D9D9"}
ROLE_FILLS = {"trigger": "FFC7CE", "quoted_false": "FFEB9C",
              "true_side": "C6EFCE", "unverified": "D9D9D9"}


def to_binary(label):
    """Collapse any model's label (6-class, Model 1 3-class, or roll-up) to
    misinfo / ok / abstain, matching the regression-step collapse."""
    l = (label or "").strip().lower()
    if l in FALSE_SIDE or l == "misinformation":
        return "misinfo"
    if l in TRUE_SIDE:
        return "ok"
    return "abstain"            # unverified, no_checkworthy_claim, blank


def claim_role(verdict, stance):
    """How a single claim relates to its post's label, given stance gating."""
    v = (verdict or "").strip().lower()
    s = (stance or "").strip().lower()
    if v in FALSE_SIDE:
        return "trigger" if s == "asserted" else "quoted_false"
    if v in TRUE_SIDE:
        return "true_side"
    return "unverified"


def asserted_falseside_count(claims):
    return sum(1 for c in claims
               if str(c.get("verdict", "")).strip().lower() in FALSE_SIDE
               and str(c.get("stance", "")).strip().lower() == "asserted")


def _sanitize_df(df):
    """Strip control characters Excel cannot store (e.g. vertical tab, NUL) from
    every text cell. Hive post text occasionally carries these and openpyxl
    refuses to write them. Applied to every column (non-strings pass straight
    through), so it works regardless of how pandas labels string columns."""
    for col in df.columns:
        df[col] = df[col].map(
            lambda v: ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v)
    return df


def load_factcheck(path):
    if not os.path.exists(path):
        return {}, {}, {}
    data = json.load(open(path, encoding="utf-8"))
    posts = {p["doc_id"]: p for p in data.get("posts", [])}
    docs = {d["doc_id"]: d for d in data.get("documents", [])}
    by_doc = {}
    for c in data.get("claims", []):
        by_doc.setdefault(c["doc_id"], []).append(c)
    return posts, docs, by_doc


def load_engagement():
    if not os.path.exists(PARQUET):
        print(f"(No {PARQUET}; engagement columns will be blank.)")
        return {}, []
    df = pd.read_parquet(PARQUET)
    cols = [c for c in ENGAGEMENT_COLS if c in df.columns]
    if "doc_id" not in df.columns:
        df["doc_id"] = df["author"].astype(str) + "/" + df["permlink"].astype(str)
    eng = {}
    for r in df.itertuples(index=False):
        eng[getattr(r, "doc_id")] = {c: getattr(r, c, "") for c in cols}
    return eng, cols


def main():
    if not os.path.exists(M3_JSON):
        print(f"Missing {M3_JSON}. Run the Model 3 fact-check first.")
        return
    if not os.path.exists(M4_CSV):
        print(f"Missing {M4_CSV}. Run hive_build_model4.py first.")
        return

    m3_posts, m3_docs, m3_claims_by_doc = load_factcheck(M3_JSON)
    m1_posts, m1_docs, m1_claims_by_doc = load_factcheck(M1_JSON)
    m4_df = pd.read_csv(M4_CSV, dtype=str).fillna("")
    m4 = {str(r["doc_id"]).strip(): {"final_label": str(r.get("final_label", "")).strip(),
                                     "label_source": str(r.get("label_source", "")).strip()}
          for _, r in m4_df.iterrows()}
    eng, eng_cols = load_engagement()

    # ---------------- by_post ----------------
    rows = []
    for doc_id, post in m3_posts.items():
        m3_label = post.get("final_label", "")
        m3_ord = post.get("accuracy_ordinal", "")
        m1_label = m1_posts.get(doc_id, {}).get("final_label", "(no title)")
        m4_label = m4.get(doc_id, {}).get("final_label", "")
        source = m4.get(doc_id, {}).get("label_source", "")

        m3_bin, m1_bin = to_binary(m3_label), to_binary(m1_label)
        disagree = (m3_bin != m1_bin and "abstain" not in (m3_bin, m1_bin)
                    and m1_label != "(no title)")

        m3_cl = m3_claims_by_doc.get(doc_id, [])
        m1_cl = m1_claims_by_doc.get(doc_id, [])
        m1_first = m1_cl[0] if m1_cl else {}
        body = (m3_docs.get(doc_id, {}).get("clean_text", "") or "")[:BODY_CHARS]

        row = {
            "doc_id": doc_id,
            "author": post.get("author", ""),
            "title": post.get("title", ""),
            "date": post.get("date", ""),
            "category": post.get("category", ""),
            "is_political": post.get("is_political", ""),
            "sample_stratum": post.get("sample_stratum", ""),
        }
        for c in eng_cols:
            row[c] = eng.get(doc_id, {}).get(c, "")
        row.update({
            "model4_label": m4_label,
            "model4_binary": to_binary(m4_label),
            "label_source": source,
            "model3_label": m3_label,
            "model3_ordinal": m3_ord,
            "model1_label": m1_label,
            "models_disagree": "DISAGREE" if disagree else "",
            "n_model3_claims": len(m3_cl),
            "asserted_falseside_n": asserted_falseside_count(m3_cl),
            "model3_verdict_counts": json.dumps(post.get("verdict_counts", {})),
            "model1_verdict": m1_first.get("verdict", ""),
            "model1_stance": m1_first.get("stance", ""),
            "model1_query": m1_first.get("query", ""),
            "cb_max_score": post.get("cb_max_score", ""),
            "human_label": "",
            "human_notes": "",
            "body_snippet": body,
        })
        rows.append(row)
    by_post = pd.DataFrame(rows)
    by_post["_o"] = by_post["model4_binary"].map(BINARY_ORDER).fillna(9)
    by_post = by_post.sort_values(["_o", "doc_id"]).drop(columns="_o")

    # ---------------- by_claim_model3 ----------------
    def claim_rows(claims_by_doc, posts):
        out = []
        for doc_id, claims in claims_by_doc.items():
            post = posts.get(doc_id, {})
            m4_label = m4.get(doc_id, {}).get("final_label", "")
            for c in claims:
                out.append({
                    "doc_id": doc_id, "title": c.get("title", ""),
                    "category": post.get("category", ""),
                    "is_political": post.get("is_political", ""),
                    "model4_label": m4_label,
                    "label_source": m4.get(doc_id, {}).get("label_source", ""),
                    "claim": c.get("claim", ""), "query": c.get("query", ""),
                    "verdict": c.get("verdict", ""), "stance": c.get("stance", ""),
                    "claim_role": claim_role(c.get("verdict", ""), c.get("stance", "")),
                    "justification": c.get("justification", ""),
                    "results": (c.get("formatted_results", "") or "")[:RESULTS_CHARS],
                    "human_verdict": "", "human_notes": "",
                })
        return pd.DataFrame(out)

    by_claim_m3 = claim_rows(m3_claims_by_doc, m3_posts)
    by_claim_m1 = claim_rows(m1_claims_by_doc, m1_posts)

    by_post = _sanitize_df(by_post)
    by_claim_m3 = _sanitize_df(by_claim_m3)
    by_claim_m1 = _sanitize_df(by_claim_m1)
    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as xw:
        by_post.to_excel(xw, sheet_name="by_post", index=False)
        by_claim_m3.to_excel(xw, sheet_name="by_claim_model3", index=False)
        by_claim_m1.to_excel(xw, sheet_name="by_claim_model1", index=False)
        _format(xw.book)

    print(f"Wrote {OUTPUT_XLSX}: {len(by_post)} posts, "
          f"{len(by_claim_m3)} Model 3 claims, {len(by_claim_m1)} Model 1 claims")
    print("\nModel 4 binary (by_post):")
    print(by_post["model4_binary"].value_counts().to_string())
    print("\nlabel_source (by_post):")
    print(by_post["label_source"].value_counts().to_string())
    dis = int((by_post["models_disagree"] == "DISAGREE").sum())
    print(f"\nPosts where Model 1 and Model 3 disagree (both non-abstain): {dis}")


def _format(book):
    for ws in book.worksheets:
        ws.freeze_panes = "A2"
        headers = [c.value for c in ws[1]]
        for j, h in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(j)].width = WIDTHS.get(h, 16)
            if h in WRAP_COLS:
                for cell in ws[get_column_letter(j)]:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        if "model4_binary" in headers:
            col = headers.index("model4_binary") + 1
            for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                fill = POST_FILLS.get(row[0].value)
                if fill:
                    row[0].fill = PatternFill("solid", fgColor=fill)
        if "claim_role" in headers:
            col = headers.index("claim_role") + 1
            for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                fill = ROLE_FILLS.get(row[0].value)
                if fill:
                    row[0].fill = PatternFill("solid", fgColor=fill)


if __name__ == "__main__":
    main()
