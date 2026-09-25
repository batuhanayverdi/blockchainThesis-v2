"""
Build a review workbook: predictions vs gold, with transcript / query / results
================================================================================
Joins everything into one Excel so you can read the pipeline end to end against
the gold ratings:
  Sheet "by_video" : one row per gold video -> gold label, pipeline label,
                     outcome (TP/FP/FN/TN/ABSTAIN), status (checked vs
                     gate-dropped), claim count, how many claims were the
                     misinformation trigger (false + asserted), verdict_counts,
                     and a transcript snippet.
  Sheet "by_claim" : one row per checked claim -> claim, query, search results,
                     the verdict, the stance (asserted / quoted_or_questioned /
                     unclear), justification, plus the video's gold + outcome.

This matches the current stage 4 scheme: verdicts are true/false/unverified and a
post is 'misinformation' only when a claim is BOTH false AND asserted, so the
`stance` column is what explains the label. A pipeline label of 'unverified' is
an abstention: shown as outcome ABSTAIN (grey), not treated as FP/FN/TP/TN,
matching evaluate_f1.py. Errors are highlighted: FN red, FP yellow, TP green.
Sort/filter by the outcome column to inspect which FALSE CLAIM videos were missed
(FN), and which TRUE CLAIM videos were over-flagged (FP), and why.

Reads (all local):
  pilot_outputs/stage4_factcheck.json   (claims + posts from stage 4)
  pilot_transcripts.json                 (transcript text + channel/published)
  Final Scrap - Final Dataset.xlsx       (gold: Video id + Final Rating Decision)

Run:  python build_review_excel.py
Deps: pip install pandas openpyxl
"""

import json
import pandas as pd
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter


# ============================ CONFIG ========================================
FACTCHECK_JSON = "pilot_outputs/stage4_factcheck.json"
TRANSCRIPTS_JSON = "pilot_transcripts.json"
GOLD_FILE = r"C:\Users\batuh\PycharmProjects\Thesis\fact_checked_pipeline\Final Scrap - Final Dataset.xlsx"
GOLD_SHEET = "Final Dataset"
GOLD_ID_COL = "Video id"
GOLD_LABEL_COL = "Final Rating Decision"
OUTPUT_XLSX = "pilot_review.xlsx"

TRANSCRIPT_CHARS = 2000          # snippet length (Excel cells cap near 32k chars)
RESULTS_CHARS = 8000

GOLD_BIN = {"FALSE CLAIM": "misinfo", "TRUE CLAIM": "ok"}   # keys upper-cased
PRED_BIN = {"misinformation": "misinfo", "true": "ok",
            "unverified": "abstain", "no_checkworthy_claim": "ok"}
MISSING_LABEL = "(no prediction: dropped at gate / not run)"

WRAP_COLS = {"title", "transcript_snippet", "results", "justification",
             "claim", "query"}
WIDTHS = {"title": 40, "transcript_snippet": 70, "results": 70,
          "justification": 50, "claim": 50, "query": 40, "channel": 22,
          "published": 12, "verdict_counts": 22, "verdict": 12, "stance": 20,
          "asserted_false_n": 14}
FILLS = {"FN": "FFC7CE", "FP": "FFEB9C", "TP": "C6EFCE",   # red / yellow / green
         "ABSTAIN": "D9D9D9"}                              # grey


def load_gold():
    df = pd.read_excel(GOLD_FILE, sheet_name=GOLD_SHEET)
    df.columns = [str(c).strip() for c in df.columns]
    gid = GOLD_ID_COL if GOLD_ID_COL in df.columns else df.columns[1]
    glab = GOLD_LABEL_COL if GOLD_LABEL_COL in df.columns else df.columns[10]
    out = {}
    for _, r in df.iterrows():
        vid = str(r[gid]).strip()
        if vid and vid.lower() != "nan":
            out[vid] = str(r[glab]).strip().upper()       # case-robust
    return out


def outcome(gold_bin, pred_bin):
    if pred_bin == "abstain":
        return "ABSTAIN"
    if gold_bin == "misinfo":
        return "TP" if pred_bin == "misinfo" else "FN"
    return "FP" if pred_bin == "misinfo" else "TN"


def asserted_false_count(claims):
    return sum(1 for c in claims
               if str(c.get("verdict", "")).lower() == "false"
               and str(c.get("stance", "")).lower() == "asserted")


def main():
    fc = json.load(open(FACTCHECK_JSON, encoding="utf-8"))
    posts = {p["doc_id"]: p for p in fc.get("posts", [])}
    claims = fc.get("claims", [])
    claims_by_doc = {}
    for c in claims:
        claims_by_doc.setdefault(c["doc_id"], []).append(c)

    tr = {str(t.get("video_id", "")).strip(): t
          for t in json.load(open(TRANSCRIPTS_JSON, encoding="utf-8"))}
    gold = load_gold()

    # ---- by_video (one row per gold video) ----
    vrows = []
    for vid, gl in gold.items():
        meta = tr.get(vid, {})
        post = posts.get(vid)
        if post:
            plabel, status = post.get("final_label", ""), "checked"
        elif vid in tr:
            plabel, status = MISSING_LABEL, "gate-dropped"
        else:
            plabel, status = MISSING_LABEL, "no transcript"
        gb = GOLD_BIN.get(gl, "?")
        pb = PRED_BIN.get(plabel, "ok")
        vrows.append({
            "video_id": vid,
            "title": meta.get("title") or (post or {}).get("title", ""),
            "channel": meta.get("channel") or (post or {}).get("author", ""),
            "published": meta.get("published") or (post or {}).get("date", ""),
            "gold": gl, "gold_binary": gb,
            "pipeline_label": plabel, "pred_binary": pb,
            "outcome": outcome(gb, pb), "status": status,
            "n_claims": len(claims_by_doc.get(vid, [])),
            "asserted_false_n": asserted_false_count(claims_by_doc.get(vid, [])),
            "verdict_counts": json.dumps((post or {}).get("verdict_counts", {})),
            "transcript_snippet": (meta.get("transcript", "") or "")[:TRANSCRIPT_CHARS],
        })
    by_video = pd.DataFrame(vrows).sort_values(["outcome", "video_id"])

    # ---- by_claim (one row per checked claim) ----
    crows = []
    for c in claims:
        vid = c["doc_id"]
        gl = gold.get(vid, "(NOT IN GOLD)")
        gb = GOLD_BIN.get(gl, "?")
        plabel = posts.get(vid, {}).get("final_label", "")
        crows.append({
            "video_id": vid, "title": c.get("title", ""),
            "gold": gl, "gold_binary": gb, "pipeline_label": plabel,
            "video_outcome": outcome(gb, PRED_BIN.get(plabel, "ok")),
            "claim": c.get("claim", ""), "query": c.get("query", ""),
            "verdict": c.get("verdict", ""), "stance": c.get("stance", ""),
            "justification": c.get("justification", ""),
            "results": (c.get("formatted_results", "") or "")[:RESULTS_CHARS],
        })
    by_claim = pd.DataFrame(crows)

    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as xw:
        by_video.to_excel(xw, sheet_name="by_video", index=False)
        by_claim.to_excel(xw, sheet_name="by_claim", index=False)
        _format(xw.book)

    print(f"Wrote {OUTPUT_XLSX}: {len(by_video)} videos, {len(by_claim)} claims")
    print("\noutcome counts (by_video):")
    print(by_video["outcome"].value_counts().to_string())
    print("\nstatus counts:")
    print(by_video["status"].value_counts().to_string())


def _format(book):
    for ws in book.worksheets:
        ws.freeze_panes = "A2"
        headers = [c.value for c in ws[1]]
        for j, h in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(j)].width = WIDTHS.get(h, 16)
            if h in WRAP_COLS:
                for cell in ws[get_column_letter(j)]:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        oc_name = "outcome" if "outcome" in headers else (
            "video_outcome" if "video_outcome" in headers else None)
        if oc_name:
            oc = headers.index(oc_name) + 1
            for row in ws.iter_rows(min_row=2, min_col=oc, max_col=oc):
                fill = FILLS.get(row[0].value)
                if fill:
                    row[0].fill = PatternFill("solid", fgColor=fill)


if __name__ == "__main__":
    main()
